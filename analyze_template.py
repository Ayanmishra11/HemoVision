"""Analyze the SIH template PPTX structure."""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.util import Inches, Pt, Emu

prs = Presentation(r'C:\Users\VICTUS\HemoVision\HHTH1.pptx')
print(f'Slide width: {prs.slide_width}, height: {prs.slide_height}')
print(f'Slide count: {len(prs.slides)}')
print(f'Layouts available: {len(prs.slide_layouts)}')
for i, layout in enumerate(prs.slide_layouts):
    print(f'  Layout {i}: {layout.name}')
print()

for si, slide in enumerate(prs.slides):
    print(f'=== SLIDE {si+1} (layout: {slide.slide_layout.name}) ===')
    for shape in slide.shapes:
        info = f'  [{shape.shape_type}] name="{shape.name}" pos=({shape.left},{shape.top}) size=({shape.width},{shape.height})'
        if shape.has_text_frame:
            texts = []
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    texts.append(t[:100])
            if texts:
                joined = " | ".join(texts[:4])
                info += f' text: {joined}'
                if len(texts) > 4:
                    info += f' ... (+{len(texts)-4} more)'
        print(info)
    print()
