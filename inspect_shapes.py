import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation(r'C:\Users\VICTUS\HemoVision\HHTH1.pptx')
for i, slide in enumerate(prs.slides):
    print(f'=== SLIDE {i+1} ===')
    for s in slide.shapes:
        txt = ""
        if s.has_text_frame:
            txt = " ".join([p.text.strip() for p in s.text_frame.paragraphs if p.text.strip()])[:60]
        print(f'  [{s.name}] type={s.shape_type} pos=({s.left},{s.top}) size=({s.width},{s.height}) txt="{txt}"')
